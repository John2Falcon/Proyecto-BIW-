import os
import datetime
# Importamos openpyxl para xlsx
from openpyxl import load_workbook
from bs4 import BeautifulSoup

from pypdf import PdfReader
from docx import Document
from pptx import Presentation 
from clasificador_documentos import Clasificador

API_BASE_URL = "http://localhost:8000"

class ProcesadorDocumentos:

    def __init__(self):
        self.clasificador = Clasificador()
        
    def procesar_archivo(self, ruta_archivo, nombre_original):
        """
        Método Principal para archivos descargados (upload o crawler).
        """
        extension_con_punto = os.path.splitext(nombre_original)[1].lower()
        extension = extension_con_punto.replace(".", "")
        contenido = ""

        # 1. Estrategia de Extracción
        if extension_con_punto == '.pdf':
            contenido = self._extraer_pdf(ruta_archivo)
        elif extension_con_punto == '.docx':
            contenido = self._extraer_docx(ruta_archivo)
        elif extension_con_punto == '.pptx':
            contenido = self._extraer_pptx(ruta_archivo)
        elif extension_con_punto == '.xlsx': # <--- AÑADIDO: Soporte XLSX
            contenido = self._extraer_xlsx(ruta_archivo)
        elif extension_con_punto == '.txt':
            contenido = self._extraer_txt(ruta_archivo)
        else:
            raise ValueError(f"El formato '{extension_con_punto}' no está soportado por el sistema.")

        # 2. Validación de Contenido Vacío
        if not contenido or not contenido.strip():
            raise ValueError(f"El archivo '{nombre_original}' no contiene texto extraíble (¿Está vacío?).")

        # 3. Construcción
        return self._construir_doc_solr(
            titulo=nombre_original,
            contenido=contenido,
            url_id=nombre_original,
            url_visual=f"{API_BASE_URL}/descargas/{nombre_original}",
            extension=extension
        )
    
    def procesar_html_descargado(self, html_content, url_origen):
        """
        Recibe HTML crudo del Crawler y devuelve el dict para Solr.
        """
        try:
            soup = BeautifulSoup(html_content, 'html.parser')
            texto_limpio = self._limpiar_html(soup)
            
            if not texto_limpio.strip():
                return None
                
            titulo = soup.title.string if soup.title else url_origen

            return self._construir_doc_solr(
                titulo=titulo,
                contenido=texto_limpio,
                url_id=url_origen,
                url_visual=url_origen,
                extension="html"
            )
        except:
            return None

    # ==========================================
    # métodos de procesamiento de documentos
    # ==========================================

    def _extraer_pdf(self, ruta):
        texto = []
        reader = PdfReader(ruta)
        
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except:
                raise Exception("El PDF está protegido con contraseña.")

        for page in reader.pages:
            extracto = page.extract_text()
            if extracto: texto.append(extracto)
            
        return "\n".join(texto)

    def _extraer_docx(self, ruta):
        doc = Document(ruta)
        return "\n".join([p.text for p in doc.paragraphs])

    def _extraer_pptx(self, ruta):
        texto = []
        prs = Presentation(ruta)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    texto.append(shape.text)
        return "\n".join(texto)
    
    def _extraer_xlsx(self, ruta):
        """Extrae texto de archivos .xlsx usando openpyxl (Clean Code)."""
        texto = []
        # data_only=True es CRÍTICO para indexación (extrae valores, no fórmulas)
        wb = load_workbook(ruta, data_only=True)
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows(values_only=True):
                # Limpieza: Convertir a string y eliminar None
                clean_row = [str(cell) for cell in row if cell is not None]
                if clean_row:
                    texto.append(" ".join(clean_row))
        return "\n".join(texto)

    def _extraer_txt(self, ruta):
        with open(ruta, 'r', encoding='utf-8') as f:
            return f.read()

    # ==========================================
    # AUXILIARES
    # ==========================================

    def _construir_doc_solr(self, titulo, contenido, url_id, url_visual, extension):
        contenido = contenido.strip()
        
        try:
            categorias_detectadas = self.clasificador.predecir_categoria(contenido, titulo)
        except Exception as e:
            print(f"⚠️ Error en clasificación automática: {e}")
            categorias_detectadas = ["General"]

        fecha_hoy = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

        return {
            "id": url_id,
            "titulo": titulo,
            "contenido": contenido,
            "url": url_visual,
            "extension": extension,
            "categoria": categorias_detectadas, 
            "fecha": fecha_hoy,
            "importancia": 1.0
        }
    
    def _limpiar_html(self, soup):
        for tag in soup(["script", "style", "nav", "footer", "iframe", "noscript"]):
            tag.extract()
        return soup.get_text(" ", strip=True)